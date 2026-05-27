"""Seed script: load candidates.json + positions.json + jobs.xlsx → Postgres.

Run inside Docker:
    docker compose exec api python scripts/seed.py

Run locally (dev):
    DATABASE_URL=postgresql+asyncpg://hellio:pass@localhost/hellio \
    python api/scripts/seed.py
"""
from __future__ import annotations

import asyncio
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from passlib.context import CryptContext
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

# ---------------------------------------------------------------------------
# Resolve paths — Docker vs local dev
# ---------------------------------------------------------------------------

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent   # ABC/

def _resolve(docker_path: str, local_rel: str) -> Path:
    p = Path(docker_path)
    if p.exists():
        return p
    fallback = _REPO_ROOT / local_rel
    if fallback.exists():
        return fallback
    raise FileNotFoundError(f"Cannot find data file: tried {p} and {fallback}")

CANDIDATES_JSON = _resolve("/data/seed/candidates.json", "src/data/candidates.json")
POSITIONS_JSON  = _resolve("/data/seed/positions.json",  "src/data/positions.json")
JOBS_XLSX       = _resolve("/data/cvsjobs/jobs/jobs.xlsx", "CVsJobs/jobs/jobs.xlsx")

DATABASE_URL = os.environ.get(
    "DATABASE_URL",
    "postgresql+asyncpg://hellio:hellio@localhost/hellio",
)

pwd_ctx = CryptContext(schemes=["bcrypt"], deprecated="auto")

# ---------------------------------------------------------------------------
# Seed data
# ---------------------------------------------------------------------------

SEED_USERS = [
    {"email": "admin@hellio.com",     "password": "admin123",  "role": "admin"},
    {"email": "recruiter@hellio.com", "password": "recruit123", "role": "recruiter"},
    {"email": "viewer@hellio.com",    "password": "view123",   "role": "viewer"},
]

_NOW = datetime.now(timezone.utc)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _job_id(num: int) -> str:
    return f"job_{num:03d}"


async def _upsert(session: AsyncSession, table: str, rows: list[dict], conflict_col: str) -> int:
    """INSERT rows, skip on conflict. Returns insert count."""
    if not rows:
        return 0
    inserted = 0
    for row in rows:
        cols = ", ".join(row.keys())
        placeholders = ", ".join(f":{k}" for k in row.keys())
        sql = text(
            f"INSERT INTO {table} ({cols}) VALUES ({placeholders}) "
            f"ON CONFLICT ({conflict_col}) DO NOTHING"
        )
        result = await session.execute(sql, row)
        inserted += result.rowcount
    return inserted


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

async def seed_users(session: AsyncSession) -> None:
    rows = [
        {
            "email": u["email"],
            "password_hash": pwd_ctx.hash(u["password"]),
            "role": u["role"],
            "created_at": _NOW,
        }
        for u in SEED_USERS
    ]
    n = await _upsert(session, "users", rows, "email")
    print(f"  users:                  {n:3d} inserted")


async def seed_candidates(session: AsyncSession) -> dict[str, str]:
    """Insert candidates + sub-tables. Returns {full_name: id} lookup."""
    with open(CANDIDATES_JSON) as f:
        data = json.load(f)

    name_to_id: dict[str, str] = {}
    inserted_count = 0

    for c in data:
        cid = c["id"]
        contact = c.get("contact") or {}
        src = c.get("sourceCv") or {}
        name_to_id[c["fullName"]] = cid

        n = await _upsert(session, "candidates", [{
            "id":                   cid,
            "full_name":            c["fullName"],
            "headline":             c.get("headline"),
            "status":               c.get("status", "Active"),
            "email":                contact.get("email"),
            "phone":                contact.get("phone"),
            "city":                 contact.get("city"),
            "linkedin_url":         contact.get("linkedinUrl"),
            "github_url":           contact.get("githubUrl"),
            "summary":              c.get("summary"),
            "source_cv_filename":   src.get("fileName"),
            "source_cv_format":     src.get("format"),
            "source_cv_path":       src.get("path"),
        }], "id")
        inserted_count += n
        if n == 0:
            continue  # already seeded — skip sub-tables to stay idempotent

        for i, sk in enumerate(c.get("skills") or []):
            await session.execute(text(
                "INSERT INTO candidate_skills (candidate_id, name, sort_order) "
                "VALUES (:cid, :name, :sort) ON CONFLICT DO NOTHING"
            ), {"cid": cid, "name": sk["name"], "sort": i})

        for i, ex in enumerate(c.get("experience") or []):
            await session.execute(text(
                "INSERT INTO candidate_experience "
                "(candidate_id, role, company, location, start_year, end_year, highlights, sort_order) "
                "VALUES (:cid, :role, :company, :loc, :sy, :ey, :hl, :sort) "
                "ON CONFLICT DO NOTHING"
            ), {
                "cid":     cid,
                "role":    ex["role"],
                "company": ex["company"],
                "loc":     ex.get("location"),
                "sy":      ex["startYear"],
                "ey":      ex.get("endYear"),
                "hl":      ex.get("highlights") or [],
                "sort":    i,
            })

        for i, ed in enumerate(c.get("education") or []):
            await session.execute(text(
                "INSERT INTO candidate_education "
                "(candidate_id, degree, institution, start_year, end_year, sort_order) "
                "VALUES (:cid, :deg, :inst, :sy, :ey, :sort) ON CONFLICT DO NOTHING"
            ), {
                "cid":  cid,
                "deg":  ed["degree"],
                "inst": ed["institution"],
                "sy":   ed["startYear"],
                "ey":   ed["endYear"],
                "sort": i,
            })

        for i, cert in enumerate(c.get("certifications") or []):
            await session.execute(text(
                "INSERT INTO candidate_certifications (candidate_id, name, year, sort_order) "
                "VALUES (:cid, :name, :year, :sort) ON CONFLICT DO NOTHING"
            ), {"cid": cid, "name": cert["name"], "year": cert.get("year"), "sort": i})

        for i, lang in enumerate(c.get("languages") or []):
            await session.execute(text(
                "INSERT INTO candidate_languages (candidate_id, name, proficiency, sort_order) "
                "VALUES (:cid, :name, :prof, :sort) ON CONFLICT DO NOTHING"
            ), {"cid": cid, "name": lang["name"], "prof": lang.get("proficiency", ""), "sort": i})

    print(f"  candidates:             {inserted_count:3d} inserted (+ sub-tables)")
    return name_to_id


async def seed_positions(session: AsyncSession) -> None:
    with open(POSITIONS_JSON) as f:
        data = json.load(f)

    inserted_count = 0
    for p in data:
        pid = p["id"]
        src = p.get("sourceDocument") or {}
        n = await _upsert(session, "positions", [{
            "id":                       pid,
            "title":                    p["title"],
            "status":                   p.get("status", "Open"),
            "hiring_manager_email":     p.get("hiringManagerEmail"),
            "description":              p.get("description"),
            "location":                 p.get("location"),
            "seniority":                p.get("seniority"),
            "salary_range":             p.get("salaryRange"),
            "source_document_filename": src.get("fileName"),
            "source_document_path":     src.get("path"),
        }], "id")
        inserted_count += n
        if n == 0:
            continue  # already seeded — skip requirements to stay idempotent

        reqs = p.get("requirements") or {}
        for i, txt in enumerate(reqs.get("mustHave") or []):
            await session.execute(text(
                "INSERT INTO position_requirements (position_id, type, text, sort_order) "
                "VALUES (:pid, 'must_have', :txt, :sort) ON CONFLICT DO NOTHING"
            ), {"pid": pid, "txt": txt, "sort": i})
        for i, txt in enumerate(reqs.get("niceToHave") or []):
            await session.execute(text(
                "INSERT INTO position_requirements (position_id, type, text, sort_order) "
                "VALUES (:pid, 'nice_to_have', :txt, :sort) ON CONFLICT DO NOTHING"
            ), {"pid": pid, "txt": txt, "sort": i})

    print(f"  positions:              {inserted_count:3d} inserted (+ requirements)")


async def seed_applications(session: AsyncSession, name_to_id: dict[str, str]) -> None:
    """Parse jobs.xlsx — resolve candidate names → IDs, insert applications."""
    wb = openpyxl.load_workbook(JOBS_XLSX)
    ws = wb.active

    inserted = 0
    skipped_unknown = 0
    app_counter = 1

    for row in ws.iter_rows(min_row=2, values_only=True):
        job_num = row[0]
        if job_num is None:
            continue
        position_id = _job_id(int(job_num))

        # Columns 5–24: pairs of (Candidate N, Status N)
        for pair_start in range(4, 24, 2):
            name   = row[pair_start]
            status = row[pair_start + 1] if pair_start + 1 < len(row) else None
            if name is None:
                break

            candidate_id = name_to_id.get(name)
            if candidate_id is None:
                print(f"  WARN: unknown candidate name '{name}' in Excel row {job_num}", file=sys.stderr)
                skipped_unknown += 1
                continue

            app_id = f"app-{app_counter:03d}"
            app_counter += 1

            result = await session.execute(text(
                "INSERT INTO applications (id, candidate_id, position_id, status, created_at) "
                "VALUES (:id, :cid, :pid, :status, :created_at) "
                "ON CONFLICT (candidate_id, position_id) DO NOTHING"
            ), {
                "id":         app_id,
                "cid":        candidate_id,
                "pid":        position_id,
                "status":     status,
                "created_at": _NOW,
            })
            inserted += result.rowcount

    print(f"  applications:           {inserted:3d} inserted ({skipped_unknown} unknown names skipped)")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

async def main() -> None:
    print("Hellio HR — seed script")
    print(f"  database: {DATABASE_URL.split('@')[-1]}")
    print(f"  candidates: {CANDIDATES_JSON}")
    print(f"  positions:  {POSITIONS_JSON}")
    print(f"  jobs xlsx:  {JOBS_XLSX}")
    print()

    engine = create_async_engine(DATABASE_URL, echo=False)
    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with async_session() as session:
        async with session.begin():
            await seed_users(session)
            name_to_id = await seed_candidates(session)
            await seed_positions(session)
            await seed_applications(session, name_to_id)

    await engine.dispose()
    print("\nDone — all inserts idempotent (ON CONFLICT DO NOTHING).")


if __name__ == "__main__":
    asyncio.run(main())
