import io

import pytest
from docx import Document


def make_docx_bytes(text: str = "Jane Doe\njane@example.com\nSoftware Engineer") -> bytes:
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_empty_docx_bytes() -> bytes:
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def minimal_docx():
    return make_docx_bytes()


@pytest.fixture
def empty_docx():
    return make_empty_docx_bytes()
