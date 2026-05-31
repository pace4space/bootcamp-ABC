from __future__ import annotations

import io

from docx import Document as _DocxDocument
from pdfminer.high_level import extract_text as _pdf_extract

from .types import DocumentKind, ParseFormat, RawDocument


class ParseError(Exception):
    """Raised when a file cannot be parsed or yields no extractable text."""


def _detect_format(filename: str) -> ParseFormat:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        return ParseFormat(ext)
    except ValueError:
        raise ParseError(f"Unsupported file extension '.{ext}'")


def _extract_pdf(file_bytes: bytes) -> str:
    return _pdf_extract(io.BytesIO(file_bytes))


def _extract_docx(file_bytes: bytes) -> str:
    doc = _DocxDocument(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def parse_cv(file_bytes: bytes, filename: str) -> RawDocument:
    fmt = _detect_format(filename)
    if fmt == ParseFormat.TXT:
        raise ParseError(f"{filename}: CV must be PDF or DOCX, not TXT")
    raw_text = _extract_pdf(file_bytes) if fmt == ParseFormat.PDF else _extract_docx(file_bytes)
    if not raw_text.strip():
        raise ParseError(f"{filename}: parsed successfully but produced no extractable text")
    return RawDocument(
        filename=filename,
        format=fmt,
        kind=DocumentKind.CV,
        raw_text=raw_text,
        char_count=len(raw_text),
    )


def parse_position(file_bytes: bytes, filename: str) -> RawDocument:
    fmt = _detect_format(filename)
    if fmt != ParseFormat.TXT:
        raise ParseError(f"{filename}: position document must be TXT, not {fmt.value.upper()}")
    raw_text = file_bytes.decode("utf-8", errors="replace")
    if not raw_text.strip():
        raise ParseError(f"{filename}: parsed successfully but produced no extractable text")
    return RawDocument(
        filename=filename,
        format=fmt,
        kind=DocumentKind.POSITION,
        raw_text=raw_text,
        char_count=len(raw_text),
    )
